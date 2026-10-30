import base64
import io
import json
import unittest
from unittest.mock import AsyncMock, patch

import docx
from docx.enum.text import WD_BREAK
import fitz
import httpx

from document_processing import DocumentExtractor, TextSection, decode_text, needs_ocr, split_sections


def pdf_bytes(*pages):
    with fitz.open() as document:
        for text in pages:
            page = document.new_page()
            if text:
                page.insert_textbox(page.rect + (40, 40, -40, -40), text)
        return document.tobytes()


def scan_bytes(header=""):
    with fitz.open(stream=pdf_bytes("Scanned document: amount 123.45 PLN"), filetype="pdf") as source:
        image = source[0].get_pixmap().tobytes("png")
    with fitz.open() as document:
        page = document.new_page()
        page.insert_image(page.rect, stream=image)
        if header:
            page.insert_textbox(page.rect + (40, 40, -40, -40), header)
        return document.tobytes()


def ocr_response(text="Kwota: 123,45 zł", status="ok", response_status="completed"):
    return {
        "status": response_status,
        "output": [{"type": "message", "content": [{
            "type": "output_text", "text": json.dumps({"status": status, "text": text}),
        }]}],
    }


class DocumentProcessingTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.requests = []
        self.response = ocr_response()

        def respond(request):
            self.requests.append(request)
            return httpx.Response(200, json=self.response)

        self.client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        self.addAsyncCleanup(self.client.aclose)
        self.extractor = DocumentExtractor(self.client, "test-key")

    async def test_readable_pdf_skips_vision_and_preserves_page_numbers_after_blank_page(self):
        first = "First page information. " * 20
        last = "Third page information. " * 20
        sections = await self.extractor.extract(pdf_bytes(first, "", last), "file.pdf")
        self.assertEqual([section.page for section in sections], [1, 3])
        self.assertEqual([section.extraction_method for section in sections], ["pdf_text", "pdf_text"])
        self.assertFalse(self.requests)

    async def test_scanned_pdf_uses_real_page_image_and_embeddings_receive_only_transcription(self):
        sections = await self.extractor.extract(scan_bytes(), "file.pdf")
        self.assertEqual(sections, [TextSection("Kwota: 123,45 zł", "vision_ocr", 1)])
        request = self.requests[0]
        self.assertEqual(str(request.url), "https://api.openai.com/v1/responses")
        self.assertEqual(request.headers["Authorization"], "Bearer test-key")
        body = json.loads(request.content)
        self.assertFalse(body["store"])
        self.assertTrue(body["text"]["format"]["strict"])
        self.assertIn("Do not summarize", body["instructions"])
        image = body["input"][0]["content"][1]
        pixels = fitz.Pixmap(base64.b64decode(image["image_url"].split(",", 1)[1]))
        self.assertLessEqual(max(pixels.width, pixels.height), 3000)
        self.assertGreater(pixels.height, 1500)

    async def test_short_native_page_uses_vision(self):
        await self.extractor.extract(pdf_bytes("Short title"), "file.pdf")
        self.assertEqual(len(self.requests), 1)

    async def test_scanned_body_is_not_hidden_by_a_native_header(self):
        header = "Native header and document reference. " * 5
        self.assertGreater(len(header), 100)
        sections = await self.extractor.extract(scan_bytes(header), "file.pdf")
        self.assertEqual(sections[0].extraction_method, "vision_ocr")
        self.assertEqual(sections[0].text, "Kwota: 123,45 zł")

    async def test_pdf_with_broken_character_mapping_uses_vision(self):
        with patch("document_processing.fitz.Page.get_text", return_value="\ufffd" * 200):
            sections = await self.extractor.extract(pdf_bytes("visible words"), "file.pdf")
        self.assertEqual(sections[0].extraction_method, "vision_ocr")

    async def test_mixed_pdf_keeps_native_text_and_ocr_in_separate_pages(self):
        with fitz.open(stream=pdf_bytes("Native page text. " * 20), filetype="pdf") as document:
            with fitz.open(stream=scan_bytes(), filetype="pdf") as scan:
                document.insert_pdf(scan)
            content = document.tobytes()
        sections = await self.extractor.extract(content, "file.pdf")
        chunks = split_sections(sections)
        self.assertEqual([section.page for section in chunks], [1, 2])
        self.assertEqual(chunks[1].text, "Kwota: 123,45 zł")
        self.assertNotIn("Native page", chunks[1].text)

    async def test_missing_key_fails_scans_but_allows_native_pdf(self):
        extractor = DocumentExtractor(self.client, None)
        await extractor.extract(pdf_bytes("Readable document text. " * 20), "file.pdf")
        with self.assertRaisesRegex(ValueError, "OPENAI_API_KEY"):
            await extractor.extract(scan_bytes(), "file.pdf")
        self.assertFalse(self.requests)

    async def test_incomplete_ocr_is_never_accepted_as_full_page(self):
        self.response = ocr_response(response_status="incomplete")
        with self.assertRaisesRegex(ValueError, "did not complete on page 1"):
            await self.extractor.extract(scan_bytes(), "file.pdf")

    async def test_refusal_is_not_embedded_as_document_content(self):
        self.response = {"status": "completed", "output": [{"type": "message", "content": [
            {"type": "refusal", "refusal": "Cannot transcribe"},
        ]}]}
        with self.assertRaisesRegex(ValueError, "refused page 1"):
            await self.extractor.extract(scan_bytes(), "file.pdf")

    async def test_unreadable_or_malformed_ocr_fails_instead_of_omitting_a_page(self):
        for response in (
            ocr_response("", "unreadable"), ocr_response("", "ok"),
            ocr_response("[illegible]", "ok"), ocr_response("text", "unknown"),
            {"status": "completed", "output": []},
        ):
            with self.subTest(response=response):
                self.response = response
                with self.assertRaises(ValueError):
                    await self.extractor.extract(scan_bytes(), "file.pdf")

    async def test_page_with_no_text_is_skipped_without_renumbering_other_pages(self):
        self.response = ocr_response("", "no_text")
        with fitz.open(stream=scan_bytes(), filetype="pdf") as document:
            with fitz.open(stream=pdf_bytes("Readable second page. " * 20), filetype="pdf") as second:
                document.insert_pdf(second)
            content = document.tobytes()
        sections = await self.extractor.extract(content, "file.pdf")
        self.assertEqual([section.page for section in sections], [2])

    async def test_transient_vision_errors_retry_but_authentication_errors_do_not(self):
        for status, attempts in ((429, 3), (503, 3), (401, 1)):
            with self.subTest(status=status):
                client = AsyncMock()
                client.post.return_value = httpx.Response(status, request=httpx.Request("POST", "https://api.openai.com/v1/responses"))
                extractor = DocumentExtractor(client, "test-key")
                with patch("document_processing.asyncio.sleep", new=AsyncMock()):
                    with self.assertRaises(httpx.HTTPStatusError):
                        await extractor.ocr_page(b"image", 1)
                self.assertEqual(client.post.await_count, attempts)

    async def test_vision_timeout_is_retried_and_can_recover(self):
        client = AsyncMock()
        client.post.side_effect = [
            httpx.ReadTimeout("timed out"),
            httpx.Response(200, json=ocr_response(), request=httpx.Request("POST", "https://api.openai.com/v1/responses")),
        ]
        with patch("document_processing.asyncio.sleep", new=AsyncMock()):
            text = await DocumentExtractor(client, "test-key").ocr_page(b"image", 1)
        self.assertEqual(text, "Kwota: 123,45 zł")
        self.assertEqual(client.post.await_count, 2)

    async def test_encrypted_pdf_is_rejected(self):
        with fitz.open(stream=pdf_bytes("secret"), filetype="pdf") as document:
            content = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
        with self.assertRaisesRegex(ValueError, "Password-protected"):
            await self.extractor.extract(content, "file.pdf")

    async def test_docx_preserves_paragraph_table_order_nested_tables_and_headers(self):
        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "Header text"
        document.sections[0].footer.paragraphs[0].text = "Footer text"
        document.add_paragraph("Before table")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Product", "Price"
        table.cell(1, 0).text, table.cell(1, 1).text = "Coffee", "12 PLN"
        table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested detail"
        document.add_paragraph("After table")
        stream = io.BytesIO()
        document.save(stream)
        sections = await self.extractor.extract(stream.getvalue(), "file.docx")
        text = sections[0].text
        self.assertIn("Header text", text)
        self.assertIn("Footer text", text)
        self.assertIn("Product | Price", text)
        self.assertIn("Coffee\nNested detail | 12 PLN", text)
        self.assertLess(text.index("Before table"), text.index("Product"))
        self.assertLess(text.index("12 PLN"), text.index("After table"))
        self.assertIsNone(sections[0].page)

    async def test_explicit_docx_page_break_is_a_hard_chunk_boundary(self):
        document = docx.Document()
        paragraph = document.add_paragraph("First part")
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        paragraph.add_run("Second part")
        stream = io.BytesIO()
        document.save(stream)
        sections = await self.extractor.extract(stream.getvalue(), "file.docx")
        self.assertEqual([part.text for part in split_sections(sections)], ["First part", "Second part"])

    async def test_file_signatures_override_generic_mime_and_extension(self):
        sections = await self.extractor.extract(pdf_bytes("Readable document. " * 20), "download", "application/octet-stream")
        self.assertEqual(sections[0].extraction_method, "pdf_text")
        document = docx.Document()
        document.add_paragraph("DOCX content")
        stream = io.BytesIO()
        document.save(stream)
        sections = await self.extractor.extract(stream.getvalue(), "download", "application/octet-stream")
        self.assertEqual(sections[0].text, "DOCX content")

    async def test_text_form_feeds_preserve_boundaries(self):
        sections = await self.extractor.extract(b"First part\fSecond part", "file.txt")
        self.assertEqual([part.text for part in split_sections(sections)], ["First part", "Second part"])

    async def test_empty_pdf_and_unsupported_file_fail(self):
        for content, filename in ((pdf_bytes(""), "empty.pdf"), (b"", "empty.txt"), (b"binary", "file.exe")):
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                await self.extractor.extract(content, filename)


class TextProcessingTests(unittest.TestCase):
    def test_unicode_boms_and_declared_legacy_encoding(self):
        text = "Zażółć gęślą jaźń. Łódź — cena: 123 zł."
        for encoding in ("utf-8-sig", "utf-16", "utf-32"):
            with self.subTest(encoding=encoding):
                self.assertEqual(decode_text(text.encode(encoding), "text/plain"), text)
        self.assertEqual(decode_text(text.encode("cp1250"), "text/plain; charset=windows-1250"), text)

    def test_undeclared_legacy_encoding_is_detected_without_replacement_characters(self):
        text = "Zażółć gęślą jaźń. Faktura dla przedsiębiorstwa w Łodzi. Wartość zamówienia: sto złotych. " * 5
        self.assertEqual(decode_text(text.encode("cp1250"), "text/plain"), text)

    def test_quality_heuristic_catches_sparse_and_garbled_text_but_accepts_polish(self):
        for text in ("short", "\ufffd" * 200, "\ue000" * 200, "(cid:123)" * 20, "a " * 100, "!@#$" * 100):
            with self.subTest(text=text[:20]):
                self.assertTrue(needs_ocr(text))
        self.assertFalse(needs_ocr("Zażółć gęślą jaźń, wartość faktury: 123,45 zł. " * 20))

    def test_chunk_overlap_never_crosses_pages_and_retains_all_words(self):
        first = "\n\n".join(f"FIRST{i}: details about item {i}." for i in range(120))
        second = "\n\n".join(f"SECOND{i}: details about item {i}." for i in range(120))
        chunks = split_sections([TextSection(first, "pdf_text", 1), TextSection(second, "vision_ocr", 2)])
        self.assertGreater(len(chunks), 2)
        for chunk in chunks:
            self.assertLessEqual(len(chunk.text), 1000)
            self.assertNotIn("SECOND" if chunk.page == 1 else "FIRST", chunk.text)
            self.assertEqual(chunk.extraction_method, "pdf_text" if chunk.page == 1 else "vision_ocr")
        for page, marker in ((1, "FIRST"), (2, "SECOND")):
            joined = "\n".join(chunk.text for chunk in chunks if chunk.page == page)
            for i in range(120):
                self.assertIn(f"{marker}{i}:", joined)


if __name__ == "__main__":
    unittest.main()
